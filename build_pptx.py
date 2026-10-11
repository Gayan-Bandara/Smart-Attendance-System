import base64
import json
import requests
import io
from pptx import Presentation
from pptx.util import Inches, Pt

def get_mermaid_image(code):
    data = {
        "code": code,
        "mermaid": '{"theme": "default"}'
    }
    json_str = json.dumps(data)
    encoded = base64.urlsafe_b64encode(json_str.encode('utf-8')).decode('utf-8')
    url = f"https://mermaid.ink/img/{encoded}"
    
    print(f"Fetching from {url}...")
    response = requests.get(url)
    if response.status_code == 200:
        return io.BytesIO(response.content)
    else:
        print(f"Failed to fetch image: {response.status_code}")
        return None

diag1 = """flowchart LR
    subgraph Hardware ["Hardware Terminal (Edge Node)"]
        direction TB
        TokenGen["Token Generator"] --> ESP32["ESP32 Microcontroller"]
        ESP32 --> Display["TFT Display"]
    end
    
    subgraph Smartphone ["Student Smartphone"]
        direction TB
        Camera["Camera Scanner"] --> WebApp["Web Application"]
    end
    
    subgraph Cloud ["Cloud Infrastructure"]
        direction TB
        TB_Cloud["ThingsBoard Platform<br>(Rule Engine & Dashboard)"]
        FB_Cloud[("Firebase Database<br>& Authentication")]
    end
    
    ESP32 -- "1. MQTT (Push active_token)" --> TB_Cloud
    Display -. "2. Dynamic QR Code" .-> Camera
    WebApp -- "3. POST Telemetry (no-cors)" --> TB_Cloud
    TB_Cloud -- "4. Verification Result" --> FB_Cloud
    FB_Cloud -- "5. Success Polling" --> WebApp"""

diag2 = """flowchart TD
    Start(["Power ON"]) --> WiFi["Connect to WiFi & MQTT"]
    WiFi --> GenToken["Generate Random 8-digit Token"]
    GenToken --> DrawQR["Render Token as QR Code on TFT"]
    DrawQR --> SyncTB["Publish Token to ThingsBoard<br>Client Attributes"]
    SyncTB --> Wait["Delay 5 Seconds"]
    Wait --> GenToken"""

diag3 = """flowchart TD
    Start(["Student Opens Web App"]) --> Auth{"Logged in via Google?"}
    Auth -- "NO" --> Prompt["Prompt Google Sign-in"] --> Auth
    Auth -- "YES" --> Fetch["Fetch Courses from Firebase"]
    Fetch --> WaitScan["Wait for QR Scan"]
    WaitScan --> Scan["Student Scans QR Code"]
    Scan --> POST["Send telemetry to ThingsBoard<br>(qr_token, email, course)"]
    POST --> Poll["Web App: Poll Firebase for confirmation"]
    
    subgraph TB_Logic ["ThingsBoard Rule Engine"]
        direction TB
        Recv["Receive Telemetry Message"] --> Match{"Does qr_token match<br>ESP32 active_token?"}
        Match -- "YES" --> Save["Save Timeseries (Mark Attendance)"]
        Save --> API["REST API: Push success to Firebase"]
        Match -- "NO" --> Drop["Silently Drop Message"]
    end
    
    POST -.-> Recv
    API -.-> Poll
    
    Poll --> Confirm{"Confirmation found<br>within 10 seconds?"}
    Confirm -- "YES" --> Success["Display: Success Confirmed"]
    Confirm -- "NO" --> Fail["Display: QR Code Expired"]
    
    Success --> End(["End"])
    Fail --> End"""

print("Starting PowerPoint creation...")
prs = Presentation()
# Set 16:9 layout
prs.slide_width = Inches(13.333)
prs.slide_height = Inches(7.5)

# Title Slide
title_slide_layout = prs.slide_layouts[0]
slide = prs.slides.add_slide(title_slide_layout)
title = slide.shapes.title
subtitle = slide.placeholders[1]
title.text = "Smart Attendance System"
subtitle.text = "System Architecture & Workflows"

diagrams = [
    ("1. System Architecture", diag1),
    ("2. Hardware Logic (ESP32)", diag2),
    ("3. Student Check-in Logic", diag3)
]

# Diagram Slides
blank_slide_layout = prs.slide_layouts[5] # Title Only

for title_text, code in diagrams:
    print(f"Processing {title_text}...")
    slide = prs.slides.add_slide(blank_slide_layout)
    title_shape = slide.shapes.title
    title_shape.text = title_text
    
    img_stream = get_mermaid_image(code)
    if img_stream:
        # Insert image and center it
        pic = slide.shapes.add_picture(img_stream, Inches(1), Inches(1.5), height=Inches(5.5))
        pic.left = int((prs.slide_width - pic.width) / 2)
        pic.top = int(Inches(1.5) + (Inches(5.5) - pic.height) / 2)

doc_path = "System_Diagrams.pptx"
prs.save(doc_path)
print(f"Saved successfully to {doc_path}")
