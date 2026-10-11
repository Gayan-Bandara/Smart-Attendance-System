import base64
import json
import requests
import io
from docx import Document
from docx.shared import Inches, Pt
from docx.enum.text import WD_ALIGN_PARAGRAPH

def get_mermaid_image(code):
    data = {
        "code": code,
        "mermaid": '{"theme": "default"}'
    }
    json_str = json.dumps(data)
    encoded = base64.urlsafe_b64encode(json_str.encode('utf-8')).decode('utf-8')
    url = f"https://mermaid.ink/img/{encoded}"
    
    print(f"Fetching from {url}...")
    response = requests.get(url)
    if response.status_code == 200:
        return io.BytesIO(response.content)
    else:
        print(f"Failed to fetch image: {response.status_code}")
        print(response.text)
        return None

diag1 = """flowchart LR
    subgraph Hardware ["Hardware Terminal (Edge Node)"]
        direction TB
        TokenGen["Token Generator"] --> ESP32["ESP32 Microcontroller"]
        ESP32 --> Display["TFT Display"]
    end
    
    subgraph Smartphone ["Student Smartphone"]
        direction TB
        Camera["Camera Scanner"] --> WebApp["Web Application"]
    end
    
    subgraph Cloud ["Cloud Infrastructure"]
        direction TB
        TB_Cloud["ThingsBoard Platform<br>(Rule Engine & Dashboard)"]
        FB_Cloud[("Firebase Database<br>& Authentication")]
    end
    
    ESP32 -- "1. MQTT (Push active_token)" --> TB_Cloud
    Display -. "2. Dynamic QR Code" .-> Camera
    WebApp -- "3. POST Telemetry (no-cors)" --> TB_Cloud
    TB_Cloud -- "4. Verification Result" --> FB_Cloud
    FB_Cloud -- "5. Success Polling" --> WebApp"""

diag2 = """flowchart TD
    Start(["Power ON"]) --> WiFi["Connect to WiFi & MQTT"]
    WiFi --> GenToken["Generate Random 8-digit Token"]
    GenToken --> DrawQR["Render Token as QR Code on TFT"]
    DrawQR --> SyncTB["Publish Token to ThingsBoard<br>Client Attributes"]
    SyncTB --> Wait["Delay 5 Seconds"]
    Wait --> GenToken"""

diag3 = """flowchart TD
    Start(["Student Opens Web App"]) --> Auth{"Logged in via Google?"}
    Auth -- "NO" --> Prompt["Prompt Google Sign-in"] --> Auth
    Auth -- "YES" --> Fetch["Fetch Courses from Firebase"]
    Fetch --> WaitScan["Wait for QR Scan"]
    WaitScan --> Scan["Student Scans QR Code"]
    Scan --> POST["Send telemetry to ThingsBoard<br>(qr_token, email, course)"]
    POST --> Poll["Web App: Poll Firebase for confirmation"]
    
    subgraph TB_Logic ["ThingsBoard Rule Engine"]
        direction TB
        Recv["Receive Telemetry Message"] --> Match{"Does qr_token match<br>ESP32 active_token?"}
        Match -- "YES" --> Save["Save Timeseries (Mark Attendance)"]
        Save --> API["REST API: Push success to Firebase"]
        Match -- "NO" --> Drop["Silently Drop Message"]
    end
    
    POST -.-> Recv
    API -.-> Poll
    
    Poll --> Confirm{"Confirmation found<br>within 10 seconds?"}
    Confirm -- "YES" --> Success["Display: Success Confirmed"]
    Confirm -- "NO" --> Fail["Display: QR Code Expired"]
    
    Success --> End(["End"])
    Fail --> End"""

print("Starting document creation...")
document = Document()

# Title
heading = document.add_heading('Smart Attendance System Diagrams', 0)
heading.alignment = WD_ALIGN_PARAGRAPH.CENTER

# Diagrams
diagrams = [
    ("1. System Architecture (Block Diagram)", diag1),
    ("2. Hardware Logic (ESP32 Flowchart)", diag2),
    ("3. Student Check-in Logic (System Flowchart)", diag3)
]

for title, code in diagrams:
    print(f"Processing {title}...")
    p = document.add_heading(title, level=2)
    img_stream = get_mermaid_image(code)
    if img_stream:
        p_img = document.add_paragraph()
        p_img.alignment = WD_ALIGN_PARAGRAPH.CENTER
        run = p_img.add_run()
        run.add_picture(img_stream, width=Inches(6.0))
    else:
        document.add_paragraph("Failed to load diagram.")
        
    document.add_page_break()

doc_path = "System_Diagrams.docx"
document.save(doc_path)
print(f"Saved successfully to {doc_path}")
